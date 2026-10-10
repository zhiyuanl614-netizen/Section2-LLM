#!/usr/bin/env python3
"""Generate a separate, non-destructive v1.2 text-scope workbook.

This script never reads legacy v1.0 metrics/FRDI or adds D03 to the analytic
label matrix. It requires explicit v1.2 ledger, D03-excluded Step 5 metrics,
and preliminary text-scope FRDI inputs; the output path must be new.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_LEDGER = ROOT / "04_LLM抽取与Critic审查/06_v1.2文本范围预分类/B001_v1.2_scope_aligned_50record_provisional_ledger.json"
DEFAULT_FRDI = ROOT / "06_跨文献综合与FRDI/02_FRDI计算/frdi_results_v1.2_text_scope_preliminary.json"
ACTIVE_DIMS = ["D01", "D02", "D04", "D05", "D06", "D07", "D08", "D09", "D10"]


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24476B")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 12), 55)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ledger", type=Path, default=DEFAULT_LEDGER)
    parser.add_argument("--metrics", type=Path, required=True, help="Output from step5_v1.2.1_provisional_descriptive_counts.py")
    parser.add_argument("--frdi", type=Path, default=DEFAULT_FRDI)
    parser.add_argument("--output", type=Path, required=True, help="New .xlsx output path; existing file is refused")
    args = parser.parse_args()
    ledger_path = args.ledger.expanduser().resolve()
    metrics_path = args.metrics.expanduser().resolve()
    frdi_path = args.frdi.expanduser().resolve()
    output_path = args.output.expanduser().resolve()
    for path in (ledger_path, metrics_path, frdi_path):
        if not path.is_file():
            raise FileNotFoundError(path)
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("Output must be a modern .xlsx workbook")
    if output_path.exists():
        raise FileExistsError(f"Refusing to overwrite existing workbook: {output_path}")

    ledger, metrics, frdi = load(ledger_path), load(metrics_path), load(frdi_path)
    if metrics.get("D03_included_in_findings_or_metrics") is not False:
        raise ValueError("Step 5 metrics must explicitly exclude D03 from analytic outcomes")
    if set(metrics.get("primary_label_counts", {})) != set(ACTIVE_DIMS):
        raise ValueError("Metrics must contain exactly the nine v1.2 active dimensions")
    records = ledger.get("records") or []
    if ledger.get("active_analysis_dimensions") != ACTIVE_DIMS:
        raise ValueError("Ledger active dimensions do not match the v1.2 text-scope set")
    if metrics.get("D02_PDF03_boundary_status") != "unresolved_no_label_is_treated_as_user_approved":
        raise ValueError("Step 5 metrics must preserve the unresolved PDF-03/D02 boundary")
    if "D03" in metrics["primary_label_counts"]:
        raise ValueError("D03 must not appear in primary_label_counts")
    d03_records = [r for r in records if r.get("dimension_id") == "D03"]
    if not d03_records or any(
        r.get("record_status") != "scope_status_not_a_paper_finding"
        or r.get("primary_label") != "Not_Analyzed_Out_Of_Scope"
        or r.get("evidence_ids")
        or r.get("evidence")
        for r in d03_records
    ):
        raise ValueError("D03 records must be empty-evidence scope metadata, not analytic findings")
    from collections import Counter
    d03_scope_counts = dict(Counter(r.get("primary_label") for r in d03_records))
    if metrics.get("D03_scope_status_metadata_only") != d03_scope_counts:
        raise ValueError("Step 5 D03 scope metadata does not match the ledger")
    doc_ids = ledger.get("document_ids") or sorted({r.get("doc_id") for r in records if r.get("doc_id")})
    if not doc_ids:
        raise ValueError("Ledger contains no document IDs")
    by_key = {(r.get("doc_id"), r.get("dimension_id")): r for r in records}
    if len(by_key) != len(records):
        raise ValueError("Ledger contains duplicate document/dimension records")
    if any((doc, dim) not in by_key for doc in doc_ids for dim in ACTIVE_DIMS):
        raise ValueError("Ledger is missing one or more active-dimension records")
    pdf03_d02 = by_key.get(("B001-PDF-03", "D02"))
    if pdf03_d02 is None:
        raise ValueError("Pilot ledger lacks the PDF-03/D02 boundary record")
    if pdf03_d02.get("verification_status") not in {"adjudication_required", "preliminary_agent_assisted_not_final"}:
        raise ValueError("PDF-03/D02 must remain unresolved until separately adjudicated")

    wb = Workbook()
    matrix = wb.active
    matrix.title = "表2.1_9维文本标签"
    matrix.append(["文献", "维度", "主标签（输入候选）", "核验状态", "未定边界裁定状态", "来源版本", "置信度", "evidence_id数量", "机制判断（摘要）"])
    for doc_id in doc_ids:
        for dim in ACTIVE_DIMS:
            r = by_key[(doc_id, dim)]
            adjudication_note = (
                "未裁定；候选标签不代表用户批准"
                if (doc_id, dim) == ("B001-PDF-03", "D02") else ""
            )
            matrix.append([
                doc_id, dim, r.get("primary_label"), r.get("verification_status"), adjudication_note,
                r.get("source_schema_version", r.get("schema_version")),
                r.get("confidence"), len(r.get("evidence_ids") or []), r.get("mechanism_judgment"),
            ])
    style_sheet(matrix)

    counts = wb.create_sheet("表2.2_描述频数")
    counts.append(["维度", "先导语料主标签频数", "状态说明"])
    for dim in ACTIVE_DIMS:
        frequency = metrics["primary_label_counts"][dim]
        counts.append([dim, json.dumps(frequency, ensure_ascii=False, sort_keys=True), "仅描述输入的先导 ledger；不是准确率或领域发生率"])
    counts.append(["D03", json.dumps(metrics.get("D03_scope_status_metadata_only", {}), ensure_ascii=False), "范围状态 metadata；不进入主标签频数/研究发现"])
    counts.append(["D02/PDF-03", json.dumps(metrics.get("D02_PDF03_boundary_sensitivity_if_reclassified_as_Generic_Water_Network_Only"), ensure_ascii=False), "未定稿边界的敏感性情景，不是用户/专家裁定"])
    style_sheet(counts)

    frdi_ws = wb.create_sheet("表2.3_FRDI探索性")
    frdi_ws.append(["方向", "ESI", "LCI", "TGM（打印年）", "TGM（DOI令牌年）", "FRDI等权重（打印年）", "FRDI等权重（DOI令牌年）"])
    for direction_id in ("3-1", "4-1", "5-1"):
        base_print = frdi["base_case"]["printed_year_basis"][direction_id]
        base_doi = frdi["base_case"]["DOI_token_year_sensitivity"][direction_id]
        frdi_ws.append([
            direction_id, base_print["ESI"], base_print["LCI"], base_print["TGM_year_proxy"],
            base_doi["TGM_year_proxy"], base_print["FRDI_by_weight"]["equal"],
            base_doi["FRDI_by_weight"]["equal"],
        ])
    frdi_ws.append([])
    frdi_ws.append(["解释边界", "探索性敏感性计算；不得用于删除/排序定稿研究问题，也不是领域级前沿测量。"])
    frdi_ws.append(["PDF-03边界", "当前基准和备选分数仅是情景；D02仍未裁定。"])
    style_sheet(frdi_ws)

    provenance = wb.create_sheet("来源与限制")
    provenance.append(["字段", "值"])
    for key, value in [
        ("工作簿状态", "v1.2 text-scope preliminary; five-paper pilot only"),
        ("Ledger", str(ledger_path)),
        ("Ledger SHA-256", sha(ledger_path)),
        ("Step 5 metrics", str(metrics_path)),
        ("Metrics SHA-256", sha(metrics_path)),
        ("FRDI input", str(frdi_path)),
        ("FRDI SHA-256", sha(frdi_path)),
        ("Candidate-direction origin", "The three existing directions were proposed by literature comparison in the manuscript Introduction; this pipeline tests them and does not regenerate them or perform online search."),
        ("D03 policy", "Status-only metadata; excluded from findings, analytic distributions, and quality metrics."),
        ("PDF-03/D02", "Unresolved candidate label; not user-approved and not a human adjudication."),
        ("Formula policy", "No equation transcription, derivation, or equation-type claim is included."),
        ("Review status", "No independent human gold standard or independent domain-expert sign-off is available."),
        ("Corpus limit", "Five PDFs are a pilot sample; zero labels do not establish field-wide absence."),
    ]:
        provenance.append([key, value])
    style_sheet(provenance)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        raise FileExistsError(f"Target appeared during workbook generation: {output_path}")
    wb.save(output_path)
    print(f"Wrote versioned v1.2 text-scope workbook: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
